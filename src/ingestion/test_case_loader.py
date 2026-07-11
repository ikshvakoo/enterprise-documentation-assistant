from pathlib import Path

import openpyxl

from models import SourceDocument
from text_utils import clean_text
from text_utils import join_nonempty


def load_test_case_workbook(path: Path) -> list[SourceDocument]:
    """Load every worksheet from a QA test-case workbook."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    documents: list[SourceDocument] = []
    try:
        for worksheet in workbook.worksheets:
            documents.extend(load_test_case_worksheet(path, worksheet))
    finally:
        workbook.close()
    return documents


def load_test_case_worksheet(path: Path, worksheet) -> list[SourceDocument]:
    """Convert one worksheet into a readable test-case SourceDocument."""
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []

    # The first row is treated as headers so each subsequent cell can be mapped
    # to a field such as Step, Test Data, or Expected Result.
    headers = [clean_text(value) or f"Column {index + 1}" for index, value in enumerate(rows[0])]
    lines = []
    issue_key = clean_text(worksheet.title)
    test_summary = ""
    step_number = 1

    for row in rows[1:]:
        row_map = {header: clean_text(value) for header, value in zip(headers, row)}
        if not any(row_map.values()):
            continue
        # Workbooks repeat issue key and test summary on some rows. Keep the
        # latest non-empty value as context for following step rows.
        if row_map.get("Issue Key"):
            issue_key = row_map["Issue Key"]
        if row_map.get("Test Summary"):
            test_summary = row_map["Test Summary"]

        step = row_map.get("Step", "")
        test_data = row_map.get("Test Data", "")
        expected = row_map.get("Expected Result", "")
        if not any([step, test_data, expected]):
            continue

        parts = []
        # Convert spreadsheet columns into natural language so BM25 can index
        # the test flow similarly to Jira descriptions and release-note prose.
        if step:
            parts.append(f"Step {step_number}: {step}")
        if test_data:
            parts.append(f"Test data: {test_data}")
        if expected:
            parts.append(f"Expected result: {expected}")
        lines.append(" ".join(parts))
        step_number += 1

    if not lines:
        return []

    title = test_summary or worksheet.title
    return [
        SourceDocument(
            doc_id=f"test_case:{path.name}:{worksheet.title}",
            source_type="test_case",
            title=title,
            body=join_nonempty([f"Issue Key: {issue_key}", f"Test Summary: {title}", "Test Steps:", *lines]),
            source_path=str(path),
            metadata={
                "file_name": path.name,
                "sheet": worksheet.title,
                "issue_key": issue_key,
                "test_summary": title,
            },
        )
    ]
